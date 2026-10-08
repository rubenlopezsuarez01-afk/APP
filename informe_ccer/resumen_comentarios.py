"""
Resumen mensual de los comentarios de los informes diarios CCER.

Lee todos los PDF de una carpeta (un informe por día), extrae la tabla de
comentarios (página "COMMENTS", normalmente la 2ª), clasifica cada comentario
en categorías y genera un Excel con, para cada parque, cuántos días se ha
reportado cada tipo de incidencia. Ejemplo:

    Limitaciones (19 días); Meteorológico (1 día); Subestación (SET) (1 día)

Uso:
    1) Edita la variable CARPETA (más abajo) y ejecuta el script (botón "Run"), o
    2) python resumen_comentarios.py CARPETA_PDFS [-o salida.xlsx]

Requisitos:
    pip install pdfplumber openpyxl
"""

import argparse
import re
import sys
from collections import defaultdict
from datetime import datetime
from pathlib import Path

import pdfplumber
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# >>> PON AQUÍ LA RUTA DE LA CARPETA CON LOS PDF DEL MES <<<
# (la "r" delante evita problemas con las barras "\" de Windows)
# Si ejecutas el script desde la terminal pasándole una carpeta, se usa esa.
# ---------------------------------------------------------------------------
CARPETA = r"C:\Users\TU_USUARIO\Documents\Informes CCER\Septiembre"

# Excel de salida. Déjalo vacío ("") para guardarlo dentro de la carpeta.
SALIDA = r""

# ---------------------------------------------------------------------------
# Categorías: se recorren en orden y cada fragmento del comentario se asigna a
# la PRIMERA categoría cuyo patrón coincida. Edita/añade reglas a tu gusto.
# Los patrones son expresiones regulares (sin distinguir mayúsculas).
# ---------------------------------------------------------------------------
CATEGORIAS = [
    ("Subestación (SET)", [r"\bset\b", r"substation", r"transformer", r"\bline trip"]),
    ("Limitaciones", [r"limit(ed|ation)", r"regulation", r"curtail", r"\bons\b",
                      r"market strategy", r"setpoint", r"stelae", r"wake"]),
    ("Meteorológico", [r"irradiation", r"low wind", r"high wind", r"weather", r"storm",
                       r"\bice\b", r"icing", r"snow", r"lightning", r"temperature"]),
    ("Comunicaciones", [r"communication", r"\bscada\b"]),
    ("Palas", [r"blade", r"trailing edge", r"leading edge"]),
    ("Tren de potencia", [r"gearbox", r"main bearing", r"generator", r"converter"]),
    ("Mantenimiento", [r"maintenance", r"service", r"\bmx\b", r"inspection", r"replace"]),
    ("Avería aerogenerador", [r"error", r"fault", r"troubleshoot", r"stop", r"trip",
                              r"emergency", r"trigger", r"sensor", r"pitch", r"yaw",
                              r"batter", r"not ok", r"collapsed", r"pause"]),
    ("Datos", [r"\bdata\b"]),
]
CATEGORIA_POR_DEFECTO = "Otros"

# Fragmentos informativos que no son incidencias (no se cuentan).
IGNORAR = [r"accumulated wind speed average"]

_CAT_RE = [(nombre, [re.compile(p, re.I) for p in pats]) for nombre, pats in CATEGORIAS]
_IGNORAR_RE = [re.compile(p, re.I) for p in IGNORAR]


def clasificar(comentario):
    """Devuelve el conjunto de categorías presentes en un comentario."""
    categorias = set()
    # Se trocea por ";" y por "." final de frase (sin romper decimales "4,3"
    # ni códigos tipo "Batt.contactor").
    for frag in re.split(r";|\.(?=\s|$)", comentario):
        frag = frag.strip(" .")
        if not frag or any(r.search(frag) for r in _IGNORAR_RE):
            continue
        if re.fullmatch(r"monthly:?", frag, re.I):
            continue
        for nombre, regs in _CAT_RE:
            if any(r.search(frag) for r in regs):
                categorias.add(nombre)
                break
        else:
            categorias.add(CATEGORIA_POR_DEFECTO)
    return categorias


def _limpiar(texto):
    return re.sub(r"\s+", " ", texto or "").strip()


def extraer_comentarios(ruta_pdf):
    """Devuelve (fecha, [(parque, comentario), ...]) de un informe."""
    with pdfplumber.open(ruta_pdf) as pdf:
        pagina = None
        for pg in pdf.pages:
            if re.search(r"COMMENTS", pg.extract_text() or "", re.I):
                pagina = pg
                break
        if pagina is None:
            if len(pdf.pages) < 2:
                return None, []
            pagina = pdf.pages[1]

        # El PDF lleva una capa de texto oculta (blanca y de tamaño ~0) que
        # ensucia la extracción: nos quedamos solo con el texto visible.
        pagina = pagina.filter(lambda o: o.get("object_type") != "char" or o.get("size", 0) >= 1)

        texto = pagina.extract_text() or ""
        fecha = None
        m = re.search(r"COMMENTS\s*(\d{2}/\d{2}/\d{4})", texto, re.I)
        if m:
            fecha = datetime.strptime(m.group(1), "%d/%m/%Y").date()

        filas = []
        for tabla in pagina.extract_tables():
            for fila in tabla:
                celdas = [_limpiar(c) for c in fila if c is not None]
                if len(celdas) < 2:
                    continue
                parque, comentario = celdas[0], " ".join(celdas[1:])
                if not parque or re.search(r"COMMENTS", parque, re.I):
                    continue
                filas.append((parque, comentario))

    if fecha is None:  # respaldo: fecha en el nombre del fichero (dd.mm.aaaa)
        m = re.search(r"(\d{2})[._-](\d{2})[._-](\d{4})", Path(ruta_pdf).name)
        if m:
            fecha = datetime(int(m.group(3)), int(m.group(2)), int(m.group(1))).date()
    return fecha, filas


def _dias(n):
    return f"{n} día" if n == 1 else f"{n} días"


def _ajustar_columnas(ws, anchos):
    for i, ancho in enumerate(anchos, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho


def _cabecera(ws, titulos):
    ws.append(titulos)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.freeze_panes = "B2"


def generar_excel(carpeta, salida):
    pdfs = sorted(Path(carpeta).glob("*.pdf")) + sorted(Path(carpeta).glob("*.PDF"))
    pdfs = sorted(set(pdfs))
    if not pdfs:
        sys.exit(f"No se han encontrado PDF en {carpeta}")

    # dias[parque][categoria] = {fecha, ...}
    dias = defaultdict(lambda: defaultdict(set))
    detalle = []
    orden_parques = []

    for pdf in pdfs:
        try:
            fecha, filas = extraer_comentarios(pdf)
        except Exception as e:  # PDF dañado o con otro formato
            print(f"AVISO: no se pudo leer {pdf.name}: {e}")
            continue
        if not filas:
            print(f"AVISO: sin comentarios en {pdf.name}")
            continue
        clave_dia = fecha or pdf.name
        for parque, comentario in filas:
            if parque not in orden_parques:
                orden_parques.append(parque)
            cats = clasificar(comentario)
            for cat in cats:
                dias[parque][cat].add(clave_dia)
            detalle.append((fecha, pdf.name, parque, comentario,
                            "; ".join(sorted(cats)) if cats else "Sin incidencias"))
        print(f"OK  {pdf.name}: {len(filas)} parques")

    categorias = [c for c, _ in CATEGORIAS] + [CATEGORIA_POR_DEFECTO]
    categorias = [c for c in categorias if any(c in dias[p] for p in dias)]

    wb = Workbook()

    # --- Hoja Resumen -----------------------------------------------------
    ws = wb.active
    ws.title = "Resumen"
    _cabecera(ws, ["Parque", "Resumen"] + categorias)
    for parque in orden_parques:
        cats = dias.get(parque, {})
        ordenadas = sorted(cats.items(), key=lambda kv: (-len(kv[1]), kv[0]))
        resumen = "; ".join(f"{c} ({_dias(len(f))})" for c, f in ordenadas) or "Sin incidencias"
        ws.append([parque, resumen] + [len(cats[c]) if c in cats else None for c in categorias])
    for fila in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        fila[0].alignment = Alignment(wrap_text=True, vertical="top")
    _ajustar_columnas(ws, [24, 70] + [14] * len(categorias))
    ws.auto_filter.ref = ws.dimensions

    # --- Hoja Días por categoría -----------------------------------------
    ws = wb.create_sheet("Días por categoría")
    _cabecera(ws, ["Parque", "Categoría", "Nº días", "Días"])
    for parque in orden_parques:
        for cat, fechas in sorted(dias.get(parque, {}).items(), key=lambda kv: (-len(kv[1]), kv[0])):
            txt = ", ".join(f.strftime("%d/%m") if hasattr(f, "strftime") else str(f)
                            for f in sorted(fechas, key=str))
            ws.append([parque, cat, len(fechas), txt])
    _ajustar_columnas(ws, [24, 24, 10, 90])
    ws.auto_filter.ref = ws.dimensions

    # --- Hoja Detalle -----------------------------------------------------
    ws = wb.create_sheet("Detalle")
    _cabecera(ws, ["Fecha", "Fichero", "Parque", "Comentario", "Categorías"])
    for fila in sorted(detalle, key=lambda r: (str(r[0]), r[1])):
        ws.append(list(fila))
        ws.cell(ws.max_row, 1).number_format = "DD/MM/YYYY"
        ws.cell(ws.max_row, 4).alignment = Alignment(wrap_text=True, vertical="top")
    _ajustar_columnas(ws, [12, 36, 24, 100, 40])
    ws.auto_filter.ref = ws.dimensions

    wb.save(salida)
    print(f"\nExcel generado: {salida}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("carpeta", nargs="?", default=CARPETA,
                    help="Carpeta con los PDF diarios del mes (por defecto, la variable CARPETA)")
    ap.add_argument("-o", "--salida", default=SALIDA,
                    help="Ruta del Excel de salida (por defecto, dentro de la carpeta)")
    # En Jupyter/Spyder sys.argv trae argumentos del propio kernel: se ignoran
    # y se usan las variables CARPETA y SALIDA.
    en_notebook = "ipykernel" in sys.modules or "IPython" in sys.modules
    args = ap.parse_args([] if en_notebook else None)

    carpeta = Path(args.carpeta)
    if not carpeta.is_dir():
        sys.exit(f"La carpeta no existe: {carpeta}\n"
                 "Revisa la variable CARPETA al principio del script.")
    salida = Path(args.salida) if args.salida else carpeta / "Resumen_comentarios_CCER.xlsx"
    generar_excel(carpeta, salida)


if __name__ == "__main__":
    main()
